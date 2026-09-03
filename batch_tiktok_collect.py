#!/usr/bin/env python3
from __future__ import annotations

import argparse
import copy
import hashlib
import hmac
import json
import os
import re
import ssl
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

TITLE_ROOT = Path(r"D:\图片\标题\最终标题")
LOCAL_IMAGE_ROOT = Path(r"D:\图片")
MIAOSHOU_BASE_URL = os.environ.get("MIAOSHOU_BASE_URL", "https://openapi-erp.91miaoshou.com")
OSS_BASE_URL = "https://duanhah-miaoshou-picture.oss-cn-shenzhen.aliyuncs.com"
SHOP_ID = 13781675
DEFAULT_BATCH = "万圣节测试"
DEFAULT_TEMPLATE = "大地毯"
ENDPOINTS = {
    "create_common_collect_product": "/open/v1/product/common_collect_box/common_collect_box/add_common_collect_box_detail",
    "claim_common_products_to_platform": "/open/v1/product/common_collect_box/common_collect_box/claimed",
    "search_tk_collect_products": "/open/v1/product/collect_box/tiktok/collect_box/search_collect_box_detail_list",
    "get_tk_shop_collect_item_info": "/open/v1/product/collect_box/tiktok/collect_box/get_shop_collect_item_info",
    "save_tk_shop_collect_item_info": "/open/v1/product/collect_box/tiktok/collect_box/save_shop_collect_item_info",
}
TEMPLATES = {
    "大地毯": {
        "detail_id": 3325026487,
        "group_name": "tiktok  大地毯    参数模版",
    },
    "非定制毛毯": {
        "detail_id": 3337135157,
        "group_name": "tiktok  非定制毛毯   参数模版",
    },
    "定制毛毯": {
        "detail_id": 3337135151,
        "group_name": "tiktok  定制毛毯   参数模版",
    },
}
VALID_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
MiaoshouCredentials = tuple[str, str] | None


def json_dumps(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, separators=(",", ":"))


def generate_sign(app_secret: str, path: str, timestamp: int, app_key: str, body_json: str) -> str:
    content = f"{app_secret}{path}{timestamp}{app_key}{body_json}{app_secret}"
    return hmac.new(
        app_secret.encode("utf-8"),
        content.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def default_credentials() -> tuple[str, str]:
    app_key = os.environ.get("MIAOSHOU_APP_KEY", "").strip()
    app_secret = os.environ.get("MIAOSHOU_APP_SECRET", "").strip()
    if not app_key or not app_secret:
        raise ValueError("缺少妙手 MIAOSHOU_APP_KEY / MIAOSHOU_APP_SECRET；网页账号管理里每个妙手账号需要单独配置。")
    return app_key, app_secret


def post_to_miaoshou(path: str, body: dict[str, Any], credentials: tuple[str, str]) -> dict[str, Any]:
    app_key, app_secret = credentials
    body_json = json_dumps(body)
    timestamp = int(time.time())
    sign = generate_sign(app_secret, path, timestamp, app_key, body_json)
    request = urllib.request.Request(
        f"{MIAOSHOU_BASE_URL}{path}",
        data=body_json.encode("utf-8"),
        method="POST",
        headers={
            "Content-Type": "application/json",
            "x-app-key": app_key,
            "x-timestamp": str(timestamp),
            "x-sign": sign,
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = response.read().decode("utf-8")
            return json.loads(payload) if payload else {}
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Miaoshou HTTP {exc.code}: {detail}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Miaoshou network error: {exc.reason}") from exc


def post(endpoint_key: str, body: dict[str, Any], credentials: MiaoshouCredentials = None) -> dict[str, Any]:
    response: dict[str, Any] = {}
    for attempt in range(4):
        try:
            path = ENDPOINTS[endpoint_key]
            response = post_to_miaoshou(path, body, credentials or default_credentials())
        except RuntimeError as exc:
            if "httpGatewayTimeout" not in str(exc) and "HTTP 504" not in str(exc):
                raise
            if attempt == 3:
                raise
            time.sleep(2 + attempt * 2)
            continue
        if response.get("code") != "accountApiQpsRateLimit":
            return response
        time.sleep(2 + attempt * 2)
    return response


def natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value)]


def batch_from_title_file(path: Path) -> str:
    name = path.stem
    return name.removeprefix("最终标题_")


def title_file_for_batch(batch: str) -> Path:
    return TITLE_ROOT / f"最终标题_{batch}.xlsx"


def header_key(value: Any) -> str:
    return re.sub(r"[\s_：:（）()]+", "", str(value or "").strip().lower())


def find_columns(ws: Any) -> tuple[int, int, int]:
    seq_names = {"序号", "编号", "no", "number", "id"}
    title_names = {"标题", "最终英文标题", "英文标题", "产品标题", "title", "producttitle"}
    for row in range(1, min(ws.max_row, 10) + 1):
        headers = {header_key(ws.cell(row, col).value): col for col in range(1, ws.max_column + 1)}
        seq_col = next((headers[name] for name in seq_names if name in headers), None)
        title_col = next((headers[name] for name in title_names if name in headers), None)
        if seq_col and title_col:
            return row, seq_col, title_col
    raise ValueError("Excel 表头必须包含「序号」和「标题」两列；标题列也可以叫「最终英文标题」或 title。")


def read_items(title_file: Path) -> list[dict[str, Any]]:
    wb = load_workbook(title_file, data_only=True)
    ws = wb.active
    header_row, seq_col, title_col = find_columns(ws)
    items: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        seq = ws.cell(row, seq_col).value
        title = ws.cell(row, title_col).value
        if seq in (None, "") and title in (None, ""):
            continue
        if not isinstance(seq, (int, float)) or int(seq) != seq:
            raise ValueError(f"row {row}: 序号列不是整数")
        title_text = str(title or "").strip()
        if not 25 <= len(title_text) <= 255:
            raise ValueError(f"seq {int(seq)}: 标题长度 {len(title_text)} 不在 25-255 范围")
        items.append({"seq": int(seq), "title": title_text, "excel_row": row})
    return items


def image_urls(
    batch: str,
    seq: int,
    local_image_root: Path = LOCAL_IMAGE_ROOT,
    image_prefix: str | None = None,
) -> list[str]:
    prefix = image_prefix or batch
    folder = local_image_root / str(seq)
    if folder.exists():
        url_prefix = prefix
    else:
        url_prefix = prefix
        batch_folder = local_image_root / batch
        folder = (batch_folder if batch_folder.exists() else local_image_root) / str(seq)
    if not folder.exists():
        raise FileNotFoundError(f"缺少图片文件夹: {folder}")
    files = sorted(
        [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in VALID_IMAGE_EXTS],
        key=lambda p: natural_key(p.name),
    )
    if not files:
        raise ValueError(f"图片文件夹为空: {folder}")
    if len(files) > 15:
        raise ValueError(f"seq {seq}: 图片数量 {len(files)} 超过妙手限制 15 张")
    return [
        f"{OSS_BASE_URL}/{urllib.parse.quote(f'{url_prefix}/{seq}/{file.name}', safe='/')}"
        for file in files
    ]


def check_url(url: str) -> dict[str, Any]:
    request = urllib.request.Request(url, method="HEAD")
    with urllib.request.urlopen(request, timeout=15, context=ssl.create_default_context()) as response:
        return {
            "status": response.status,
            "content_type": response.headers.get("Content-Type"),
            "content_length": int(response.headers.get("Content-Length") or 0),
        }


def find_template_detail_id(template_title: str, shop_id: int, credentials: MiaoshouCredentials = None) -> int:
    for page in range(1, 6):
        response = post(
            "search_tk_collect_products",
            {"pageNo": page, "pageSize": 100, "status": "notPublished", "title": template_title},
            credentials,
        )
        for item in (response.get("data") or {}).get("detailList") or []:
            shops = item.get("collectBoxDetailShopList") or []
            if item.get("title") == template_title and any(str(s.get("shopId")) == str(shop_id) for s in shops):
                return int(item["collectBoxDetailId"])
    raise RuntimeError(f"未找到 TikTok 模板产品: {template_title}")


def get_tk_info(detail_id: int, shop_id: int, credentials: MiaoshouCredentials = None) -> tuple[str, dict[str, Any]]:
    response = post("get_tk_shop_collect_item_info", {"detailId": detail_id, "shopId": shop_id}, credentials)
    if response.get("result") != "success":
        raise RuntimeError(f"查询 TikTok 详情失败: {response}")
    data = response.get("data") or {}
    oss_md5 = data.get("ossMd5")
    info = data.get("shopCollectItemInfo")
    if not oss_md5 or not isinstance(info, dict):
        raise RuntimeError(f"TikTok 详情缺少 ossMd5/shopCollectItemInfo: {response}")
    return str(oss_md5), info


def create_common_product(
    title: str,
    urls: list[str],
    template_info: dict[str, Any],
    item_num: str,
    credentials: MiaoshouCredentials = None,
) -> int:
    body = {
        "title": title,
        "itemNum": item_num[:50],
        "notesText": "Area rug batch product.",
        "notes": template_info.get("notes") or "<p>Area rug batch product.</p>",
        "sourceAttrs": [],
        "price": first_template_price(template_info),
        "stock": 0,
        "imgUrls": urls,
        "weight": template_info.get("weight") or 8,
        "packageLength": template_info.get("packageLength") or 50,
        "packageWidth": template_info.get("packageWidth") or 45,
        "packageHeight": template_info.get("packageHeight") or 40,
    }
    response = post("create_common_collect_product", body, credentials)
    if response.get("result") != "success":
        raise RuntimeError(f"创建公共采集箱产品失败: {response}")
    return int((response.get("data") or {})["commonCollectBoxDetailId"])


def first_template_price(template_info: dict[str, Any]) -> float:
    sku_map = template_info.get("skuMap") or {}
    for sku in sku_map.values():
        price = sku.get("priceIncludeVat") or sku.get("price") or sku.get("originPrice")
        if price:
            return float(price)
    return 79.49


def claim_to_tiktok(common_id: int, credentials: MiaoshouCredentials = None) -> int:
    body = {
        "detailSerialNumberPlatformList": [
            {"detailId": common_id, "platform": "tiktok", "serialNumber": 1}
        ]
    }
    response = post("claim_common_products_to_platform", body, credentials)
    if response.get("result") != "success":
        raise RuntimeError(f"认领到 TikTok 失败: {response}")
    mapping = ((response.get("data") or {}).get("platformCollectBoxDetailIdMap") or {}).get("tiktok") or {}
    if str(common_id) not in mapping:
        raise RuntimeError(f"认领成功但未返回 TikTok detailId: {response}")
    return int(mapping[str(common_id)])


def find_existing_by_title(title: str, shop_id: int, credentials: MiaoshouCredentials = None) -> dict[str, Any] | None:
    for page in range(1, 6):
        response = post(
            "search_tk_collect_products",
            {"pageNo": page, "pageSize": 100, "status": "notPublished", "title": title},
            credentials,
        )
        for item in (response.get("data") or {}).get("detailList") or []:
            shops = item.get("collectBoxDetailShopList") or []
            if item.get("title") == title and any(str(s.get("shopId")) == str(shop_id) for s in shops):
                return item
    return None


def save_from_template(
    detail_id: int,
    shop_id: int,
    title: str,
    urls: list[str],
    template_info: dict[str, Any],
    credentials: MiaoshouCredentials = None,
) -> dict[str, Any]:
    oss_md5, _ = get_tk_info(detail_id, shop_id, credentials)
    info = copy.deepcopy(template_info)
    info.update(
        {
            "detailId": detail_id,
            "shopId": shop_id,
            "title": title,
            "oriTitle": title,
            "imgUrls": urls,
            "deliveryOptionSetType": "default",
        }
    )
    if not info.get("sizeChart"):
        info.pop("sizeChart", None)
        info.pop("sizeChartType", None)
    for shop in info.get("collectBoxDetailShopList") or []:
        shop["shopId"] = shop_id
        shop.setdefault("site", info.get("site", "US"))
        shop.setdefault("deliveryOptionSetType", "default")
    response = post(
        "save_tk_shop_collect_item_info",
        {"ossMd5": oss_md5, "detailId": detail_id, "shopId": shop_id, "shopCollectItemInfo": info},
        credentials,
    )
    if response.get("result") != "success":
        raise RuntimeError(f"保存 TikTok 详情失败: {response}")
    return response


def verify_product(
    detail_id: int,
    shop_id: int,
    title: str,
    urls: list[str],
    template_info: dict[str, Any],
    credentials: MiaoshouCredentials = None,
) -> dict[str, Any]:
    _, info = get_tk_info(detail_id, shop_id, credentials)
    sizes = [
        value.get("attrValue")
        for prop in info.get("skuPropertyList") or []
        if prop.get("attrName") == "Size"
        for value in prop.get("attrValueList") or []
    ]
    return {
        "title_ok": info.get("title") == title,
        "images_ok": (info.get("imgUrls") or []) == urls,
        "image_count": len(info.get("imgUrls") or []),
        "cid_ok": info.get("cid") == template_info.get("cid"),
        "sku_count": len(info.get("skuMap") or {}),
        "sizes": sizes,
    }


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def previous_successes(batch: str, template_detail_id: int, log_dir: Path = Path("runs")) -> dict[int, dict[str, Any]]:
    runs = sorted(log_dir.glob(f"{batch}_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    found: dict[int, dict[str, Any]] = {}
    for path in runs:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        summary = data.get("summary") or {}
        if summary.get("templateDetailId") != template_detail_id:
            continue
        for item in data.get("results") or []:
            if item.get("status") == "success" and item.get("tiktokDetailId") and item.get("seq") not in found:
                found[int(item["seq"])] = item
    return found


def run_batch(
    *,
    batch: str,
    title_file: Path,
    local_image_root: Path = LOCAL_IMAGE_ROOT,
    shop_id: int = SHOP_ID,
    template: str = DEFAULT_TEMPLATE,
    template_title: str | None = None,
    template_detail_id: int | None = None,
    image_prefix: str | None = None,
    limit: int | None = None,
    only_seqs: list[int] | None = None,
    dry_run: bool = False,
    reuse_existing: bool = True,
    log_dir: Path = Path("runs"),
    progress: Callable[[dict[str, Any]], None] | None = None,
    miaoshou_credentials: MiaoshouCredentials = None,
) -> dict[str, Any]:
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = log_dir / f"{batch}_{run_id}.json"

    items = read_items(title_file)
    if limit:
        items = items[:limit]
    if only_seqs is not None:
        allowed = set(only_seqs)
        items = [item for item in items if int(item["seq"]) in allowed]

    if template_detail_id is None and template_title:
        template_detail_id = find_template_detail_id(template_title, shop_id, miaoshou_credentials)
    if template_detail_id is None:
        template_detail_id = TEMPLATES[template]["detail_id"]
    _, template_info = get_tk_info(template_detail_id, shop_id, miaoshou_credentials)
    previous = {} if dry_run or not reuse_existing else previous_successes(batch, template_detail_id, log_dir)

    results: list[dict[str, Any]] = []
    for item in items:
        seq = item["seq"]
        result: dict[str, Any] = {"seq": seq, "title": item["title"], "status": "started"}
        try:
            urls = image_urls(batch, seq, local_image_root, image_prefix)
            result["image_count"] = len(urls)
            result["ignored_files"] = ["Thumbs.db"]
            result["url_checks"] = [check_url(url) for url in urls]
            if dry_run:
                result["status"] = "dry_run_ok"
                results.append(result)
                if progress:
                    progress(result)
                continue

            existing = previous.get(seq)
            if existing:
                detail_id = int(existing["tiktokDetailId"])
                common_id = int(existing.get("commonCollectBoxDetailId") or 0)
                result["reused_from_log"] = True
                try:
                    get_tk_info(detail_id, shop_id, miaoshou_credentials)
                except Exception as exc:
                    result["stale_logged_detail"] = {"detailId": detail_id, "error": repr(exc)}
                    existing = None
            else:
                result["reused_from_log"] = False

            if not existing:
                searched = find_existing_by_title(item["title"], shop_id, miaoshou_credentials) if reuse_existing else None
                if searched:
                    detail_id = int(searched["collectBoxDetailId"])
                    common_id = int(searched.get("commonCollectBoxDetailId") or 0)
                    result["reused_existing"] = True
                else:
                    common_id = create_common_product(
                        item["title"],
                        urls,
                        template_info,
                        f"MSBATCH-{template}-{batch}-{seq}",
                        miaoshou_credentials,
                    )
                    detail_id = claim_to_tiktok(common_id, miaoshou_credentials)
                    result["reused_existing"] = False
                result.setdefault("reused_from_log", False)

            save_from_template(detail_id, shop_id, item["title"], urls, template_info, miaoshou_credentials)
            verification = verify_product(detail_id, shop_id, item["title"], urls, template_info, miaoshou_credentials)
            result.update(
                {
                    "status": "success" if all(
                        [
                            verification["title_ok"],
                            verification["images_ok"],
                            verification["cid_ok"],
                            verification["sku_count"] == len(template_info.get("skuMap") or {}),
                        ]
                    ) else "verify_failed",
                    "commonCollectBoxDetailId": common_id,
                    "tiktokDetailId": detail_id,
                    "shopId": shop_id,
                    "verification": verification,
                }
            )
        except Exception as exc:
            result.update({"status": "failed", "error": repr(exc)})
        results.append(result)
        if progress:
            progress(result)
        write_json(log_path, {"batch": batch, "templateDetailId": template_detail_id, "results": results})
        time.sleep(1.2)

    ok_statuses = {"success", "dry_run_ok"}
    summary = {
        "batch": batch,
        "titleFile": str(title_file),
        "imageRoot": str(local_image_root),
        "imagePrefix": image_prefix or batch,
        "template": template,
        "templateTitle": template_title,
        "templateDetailId": template_detail_id,
        "templateGroupName": TEMPLATES.get(template, {}).get("group_name"),
        "reuseExisting": reuse_existing,
        "shopId": shop_id,
        "total": len(results),
        "success": sum(1 for r in results if r["status"] == "success"),
        "dryRunOk": sum(1 for r in results if r["status"] == "dry_run_ok"),
        "failed": [r for r in results if r["status"] not in ok_statuses],
        "logPath": str(log_path.resolve()),
    }
    write_json(log_path, {"summary": summary, "results": results})
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--batch")
    parser.add_argument("--title-file")
    parser.add_argument("--image-root", default=str(LOCAL_IMAGE_ROOT))
    parser.add_argument("--shop-id", type=int, default=SHOP_ID)
    parser.add_argument("--template", choices=sorted(TEMPLATES), default=DEFAULT_TEMPLATE)
    parser.add_argument("--template-title")
    parser.add_argument("--template-detail-id", type=int)
    parser.add_argument("--image-prefix")
    parser.add_argument("--list-templates", action="store_true")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-reuse-existing", action="store_true")
    args = parser.parse_args()

    if args.list_templates:
        print(json.dumps(TEMPLATES, ensure_ascii=False, indent=2))
        return 0

    batch = args.batch or DEFAULT_BATCH
    title_file = Path(args.title_file) if args.title_file else title_file_for_batch(batch)
    if args.title_file and not args.batch:
        batch = batch_from_title_file(title_file)

    summary = run_batch(
        batch=batch,
        title_file=title_file,
        local_image_root=Path(args.image_root),
        shop_id=args.shop_id,
        template=args.template,
        template_title=args.template_title,
        template_detail_id=args.template_detail_id,
        image_prefix=args.image_prefix,
        limit=args.limit,
        dry_run=args.dry_run,
        reuse_existing=not args.no_reuse_existing,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not summary["failed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
